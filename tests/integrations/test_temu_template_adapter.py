import hashlib
import json
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from core.config import ReportSettings
from integrations.temu.template_adapter import TemuExcelTemplateAdapter
from modules.listing.schemas import (
    CanonicalListingRow,
    ListingBatch,
    ListingProduct,
    ListingSKU,
)


def test_template_adapter_preserves_sheets_format_and_writes_fields(tmp_path: Path) -> None:
    template = tmp_path / "temu-test-template.xlsx"
    workbook = Workbook()
    upload = workbook.worksheets[0]
    upload.title = "Upload"
    headers = ["Title", "SKU", "Width", "Height", "Colors", "Framed", "Main image"]
    for column, header in enumerate(headers, start=1):
        cell = upload.cell(row=3, column=column, value=header)
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    lists = workbook.create_sheet("Lists")
    lists.sheet_state = "hidden"
    lists["A1"] = "DO NOT REMOVE"
    workbook.save(template)
    template_hash = hashlib.sha256(template.read_bytes()).hexdigest()
    mapping = tmp_path / "mapping.json"
    mapping.write_text(
        json.dumps(
            {
                "platform": "TEMU_TEST_FIXTURE",
                "template_sha256": template_hash,
                "sheet_name": "Upload",
                "header_row": 3,
                "first_data_row": 4,
                "columns": {
                    "title": "A",
                    "sku": "B",
                    "width_cm": "C",
                    "height_cm": "D",
                    "colors_count": "E",
                    "framed": "F",
                    "main_image": "G",
                },
                "expected_headers": {
                    "title": "Title",
                    "sku": "SKU",
                    "width_cm": "Width",
                    "height_cm": "Height",
                    "colors_count": "Colors",
                    "framed": "Framed",
                    "main_image": "Main image",
                },
            }
        ),
        encoding="utf-8",
    )
    row = CanonicalListingRow(
        design_id="DESIGN-CAT-000001",
        factory_design_code="CAT000001",
        spu_item_code="PBN-CAT000001",
        title="Cat - Animals - Paint by Numbers Kit",
        title_zh="猫咪成人数字油画套装",
        theme="Animals",
        product_type="PAINT_BY_NUMBERS",
        target_market="US",
        main_image="/tmp/CAT000001.jpg",
        additional_images="",
        origin_country="中国大陆",
        material="油画布",
        sku="PBN-CAT000001-4050-24-U",
        width_cm=Decimal("40"),
        height_cm=Decimal("50"),
        colors_count=24,
        framed=False,
        factory_cost_cny=Decimal("44"),
        shipping_cost_usd=Decimal("4"),
        unit_variable_cost_cny=Decimal("73"),
    )
    output = tmp_path / "filled.xlsx"
    adapter = TemuExcelTemplateAdapter(ReportSettings(), Path.cwd())
    adapter.export(
        template_path=template,
        mapping_path=mapping,
        rows=[row],
        output_path=output,
        allow_test_fixture=True,
    )
    loaded = load_workbook(output)
    assert loaded.sheetnames == ["Upload", "Lists"]
    assert loaded["Lists"].sheet_state == "hidden"
    assert loaded["Lists"]["A1"].value == "DO NOT REMOVE"
    assert loaded["Upload"]["A3"].fill.fgColor.rgb.endswith("1F4E78")
    assert loaded["Upload"]["A4"].value == row.title
    assert loaded["Upload"]["B4"].value == row.sku
    assert loaded["Upload"]["F4"].value == "NO"


def test_template_adapter_writes_one_spu_row_before_sku_rows(tmp_path: Path) -> None:
    template = tmp_path / "temu-hierarchical-template.xlsx"
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "Upload"
    headers = ["Level", "SPU", "Title", "SKU", "Size"]
    for column, header in enumerate(headers, start=1):
        sheet.cell(row=1, column=column, value=header)
    workbook.save(template)
    mapping = tmp_path / "mapping.json"
    mapping.write_text(
        json.dumps(
            {
                "platform": "TEMU_TEST_FIXTURE",
                "template_sha256": hashlib.sha256(template.read_bytes()).hexdigest(),
                "sheet_name": "Upload",
                "header_row": 1,
                "first_data_row": 2,
                "row_strategy": "spu_with_sku_rows",
                "columns": {
                    "title": "C",
                    "sku": "D",
                    "width_cm": "E",
                    "height_cm": "E",
                    "colors_count": "E",
                    "framed": "E",
                    "main_image": "E",
                },
                "spu_columns": {
                    "level": {"column": "A", "value": "spu"},
                    "spu": {"column": "B", "source": "spu_item_code"},
                    "title": {"column": "C", "source": "title"},
                },
                "sku_columns": {
                    "level": {"column": "A", "value": "sku"},
                    "spu": {"column": "B", "source": "spu_item_code"},
                    "sku": {"column": "D", "source": "sku"},
                    "size": {"column": "E", "transform": "size_label"},
                },
            }
        ),
        encoding="utf-8",
    )
    base = {
        "design_id": "DESIGN-CAT-000001",
        "factory_design_code": "CAT000001",
        "spu_item_code": "PBN-CAT000001",
        "title": "Cat - Animals - Paint by Numbers Kit",
        "title_zh": "猫咪成人数字油画套装",
        "theme": "Animals",
        "product_type": "PAINT_BY_NUMBERS",
        "target_market": "US",
        "main_image": "/tmp/CAT000001.jpg",
        "additional_images": "",
        "origin_country": "中国大陆",
        "material": "油画布",
        "colors_count": 24,
        "framed": False,
        "factory_cost_cny": Decimal("44"),
        "shipping_cost_usd": Decimal("4"),
        "unit_variable_cost_cny": Decimal("73"),
    }
    rows = [
        CanonicalListingRow(
            **base,
            sku="PBN-CAT000001-3040-24-U",
            width_cm=Decimal("30"),
            height_cm=Decimal("40"),
        ),
        CanonicalListingRow(
            **base,
            sku="PBN-CAT000001-4050-24-U",
            width_cm=Decimal("40"),
            height_cm=Decimal("50"),
        ),
    ]
    output = tmp_path / "filled.xlsx"
    TemuExcelTemplateAdapter(ReportSettings(), Path.cwd()).export(
        template_path=template,
        mapping_path=mapping,
        rows=rows,
        output_path=output,
        allow_test_fixture=True,
    )

    loaded = load_workbook(output, data_only=False)
    assert loaded["Upload"]["A2"].value == "spu"
    assert loaded["Upload"]["B2"].value == "PBN-CAT000001"
    assert loaded["Upload"]["A3"].value == "sku"
    assert loaded["Upload"]["D3"].value == "PBN-CAT000001-3040-24-U"
    assert loaded["Upload"]["E4"].value == "40x50cm"


def test_semi_managed_preflight_blocks_unapproved_compliance_and_rights(
    tmp_path: Path,
) -> None:
    image = tmp_path / "image.png"
    image.write_bytes(b"\x89PNG\r\n\x1a\n" + b"\x00" * 8 + (900).to_bytes(4, "big") * 2)
    compliance = tmp_path / "compliance.json"
    compliance.write_text(
        json.dumps(
            {
                "jurisdiction": "US",
                "product_linkage_confirmed": False,
                "astm_d4236_toxicologist_review_confirmed": False,
                "approved_for_publish": False,
            }
        ),
        encoding="utf-8",
    )
    sku = ListingSKU(
        factory_sku="PBN-CAT000001-4050-24-U",
        width_cm=Decimal("40"),
        height_cm=Decimal("50"),
        colors_count=24,
        framed=False,
        sku_image_ref="TEMU_SKU_IMAGE_1",
        declared_price_cny=Decimal("60"),
        warehouse_inventory=10,
        package_length_cm=Decimal("55"),
        package_width_cm=Decimal("8"),
        package_height_cm=Decimal("5.5"),
        package_weight_g=Decimal("399"),
        factory_cost=None,
        shipping_cost=None,
        unit_cost=None,
    )
    product = ListingProduct(
        design_id="DESIGN-CAT-000001",
        factory_design_code="CAT000001",
        title="Cat Paint by Numbers Kit",
        title_zh="猫咪成人数字油画套装",
        theme="Animals",
        main_image=image,
        additional_images=[],
        publish_image_refs=["ASSET_1", "ASSET_2", "ASSET_3"],
        sensitive_attributes=["膏体"],
        compliance_manifest=compliance,
        product_type="PAINT_BY_NUMBERS",
        target_market="US",
        skus=[sku],
    )

    issues = TemuExcelTemplateAdapter(ReportSettings(), Path.cwd()).validate_listing(
        ListingBatch(products=[product]),
        {"row_strategy": "spu_with_sku_rows", "profile": {}},
    )
    codes = {issue.code for issue in issues}

    assert "ASSET_RIGHTS_UNCONFIRMED" in codes
    assert "COMPLIANCE_NOT_APPROVED" in codes
