from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from integrations.excel.reader import WorksheetData


class FactoryQuoteReadError(RuntimeError):
    pass


class FactoryQuoteExcelReader:
    def read(self, file_path: Path, sheet_name: str | None = None) -> WorksheetData:
        path = file_path.expanduser().resolve()
        if not path.is_file():
            raise FactoryQuoteReadError(f"Factory quote file does not exist: {path}")
        if path.suffix.lower() != ".xlsx":
            raise FactoryQuoteReadError("Factory quote input must be an .xlsx file")

        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except (InvalidFileException, OSError, ValueError) as exc:
            raise FactoryQuoteReadError(f"Cannot read factory quote workbook: {exc}") from exc

        try:
            if sheet_name is not None:
                if sheet_name not in workbook.sheetnames:
                    available = ", ".join(workbook.sheetnames)
                    raise FactoryQuoteReadError(
                        f"Worksheet {sheet_name!r} not found; available: {available}"
                    )
                worksheet = workbook[sheet_name]
            else:
                worksheet = workbook[workbook.sheetnames[0]]

            rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
            return WorksheetData(name=worksheet.title, rows=rows)
        finally:
            workbook.close()
