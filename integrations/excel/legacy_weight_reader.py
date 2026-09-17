import shutil
import subprocess
import tempfile
from pathlib import Path

from openpyxl import load_workbook

from integrations.excel.reader import WorksheetData


class LegacyWeightReadError(RuntimeError):
    pass


class LegacyWeightExcelReader:
    def __init__(self, soffice_executable: str = "soffice") -> None:
        self.soffice_executable = soffice_executable

    def read(self, file_path: Path, sheet_name: str | None = None) -> WorksheetData:
        source = file_path.expanduser().resolve()
        if not source.is_file() or source.suffix.lower() != ".xls":
            raise LegacyWeightReadError(f"Expected an existing .xls weight table: {source}")
        executable = shutil.which(self.soffice_executable)
        if executable is None:
            raise LegacyWeightReadError(
                "Legacy .xls parsing requires LibreOffice/soffice; no data was guessed"
            )
        with tempfile.TemporaryDirectory(prefix="temu-weight-") as temp_dir:
            result = subprocess.run(
                [
                    executable,
                    "--headless",
                    "--convert-to",
                    "xlsx",
                    "--outdir",
                    temp_dir,
                    str(source),
                ],
                check=False,
                capture_output=True,
                text=True,
            )
            converted = Path(temp_dir) / f"{source.stem}.xlsx"
            if result.returncode != 0 or not converted.is_file():
                raise LegacyWeightReadError(
                    f"LibreOffice could not convert the .xls weight table: {result.stderr.strip()}"
                )
            workbook = load_workbook(converted, read_only=True, data_only=True)
            try:
                target = sheet_name or workbook.sheetnames[0]
                if target not in workbook.sheetnames:
                    raise LegacyWeightReadError(f"Worksheet {target!r} does not exist")
                worksheet = workbook[target]
                return WorksheetData(
                    name=target,
                    rows=[list(row) for row in worksheet.iter_rows(values_only=True)],
                )
            finally:
                workbook.close()
