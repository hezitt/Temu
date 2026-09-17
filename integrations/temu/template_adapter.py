import hashlib
import json
import subprocess
import tempfile
from pathlib import Path
from typing import Any, Protocol, cast

from openpyxl import load_workbook

from core.config import ReportSettings
from modules.listing.image_validator import (
    ImageInspectionError,
    inspect_image,
    is_publishable_image_ref,
)
from modules.listing.schemas import CanonicalListingRow, ListingBatch, ListingIssue, ListingProduct


class TemuTemplateError(RuntimeError):
    pass


class TemuListingAdapter(Protocol):
    def export(
        self,
        *,
        template_path: Path,
        mapping_path: Path,
        rows: list[CanonicalListingRow],
        output_path: Path,
    ) -> Path: ...


class TemuExcelTemplateAdapter:
    REQUIRED_CANONICAL_FIELDS = {
        "title",
        "sku",
        "width_cm",
        "height_cm",
        "colors_count",
        "framed",
        "main_image",
    }

    def __init__(self, reports: ReportSettings, project_root: Path) -> None:
        self.node_executable = reports.node_executable
        self.writer_script = project_root / "scripts" / "fill_temu_template.mjs"

    def validate_mapping(
        self, template_path: Path, mapping_path: Path, *, allow_test_fixture: bool = False
    ) -> dict[str, Any]:
        template = template_path.expanduser().resolve()
        mapping_file = mapping_path.expanduser().resolve()
        if not template.is_file() or template.suffix.lower() != ".xlsx":
            raise TemuTemplateError("A real Temu .xlsx template is required")
        if not mapping_file.is_file():
            raise TemuTemplateError("A reviewed Temu field mapping JSON is required")
        try:
            raw_mapping = json.loads(mapping_file.read_text(encoding="utf-8"))
        except (json.JSONDecodeError, OSError) as exc:
            raise TemuTemplateError(f"Cannot read Temu mapping: {exc}") from exc
        if not isinstance(raw_mapping, dict):
            raise TemuTemplateError("Temu mapping must be a JSON object")
        mapping = cast(dict[str, Any], raw_mapping)
        allowed_platforms = {"TEMU", "TEMU_TEST_FIXTURE"} if allow_test_fixture else {"TEMU"}
        if mapping.get("platform") not in allowed_platforms:
            raise TemuTemplateError(
                "Mapping must explicitly identify platform=TEMU; Fangguo templates are not accepted"
            )
        columns = mapping.get("columns")
        if not isinstance(columns, dict):
            raise TemuTemplateError("Mapping columns must be an object")
        missing = self.REQUIRED_CANONICAL_FIELDS - set(columns)
        if missing:
            raise TemuTemplateError(
                "Temu mapping is missing canonical fields: " + ", ".join(sorted(missing))
            )
        expected_hash = mapping.get("template_sha256")
        actual_hash = hashlib.sha256(template.read_bytes()).hexdigest()
        if expected_hash and expected_hash != actual_hash:
            raise TemuTemplateError("Temu template SHA-256 does not match the reviewed mapping")
        sheet_name = mapping.get("sheet_name")
        if not isinstance(sheet_name, str) or not sheet_name:
            raise TemuTemplateError("Temu mapping requires sheet_name")
        anchors = mapping.get("template_anchors", {})
        if anchors:
            if not isinstance(anchors, dict):
                raise TemuTemplateError("template_anchors must be an object")
            workbook = load_workbook(template, read_only=True, data_only=False)
            try:
                if sheet_name not in workbook.sheetnames:
                    raise TemuTemplateError(f"Mapped worksheet not found: {sheet_name}")
                sheet = workbook[sheet_name]
                for cell, expected in anchors.items():
                    actual = sheet[str(cell)].value
                    if actual != expected:
                        raise TemuTemplateError(
                            f"Template anchor mismatch at {sheet_name}!{cell}: "
                            f"expected {expected!r}, got {actual!r}"
                        )
            finally:
                workbook.close()
        mapping["actual_template_sha256"] = actual_hash
        return mapping

    def validate_listing(
        self, listing: ListingBatch, mapping: dict[str, Any]
    ) -> list[ListingIssue]:
        if mapping.get("row_strategy", "flat") != "spu_with_sku_rows":
            return []
        profile = mapping.get("profile", {})
        if not isinstance(profile, dict):
            raise TemuTemplateError("Semi-managed mapping profile must be an object")
        min_dimension = int(profile.get("image_min_dimension_exclusive", 800))
        max_bytes = int(profile.get("image_max_bytes", 2 * 1024 * 1024))
        min_carousel = int(profile.get("carousel_min", 3))
        max_carousel = int(profile.get("carousel_max", 10))
        issues: list[ListingIssue] = []
        for product in listing.products:
            issues.extend(
                self._validate_semi_managed_product(
                    product,
                    min_dimension=min_dimension,
                    max_bytes=max_bytes,
                    min_carousel=min_carousel,
                    max_carousel=max_carousel,
                )
            )
        return issues

    def _validate_semi_managed_product(
        self,
        product: ListingProduct,
        *,
        min_dimension: int,
        max_bytes: int,
        min_carousel: int,
        max_carousel: int,
    ) -> list[ListingIssue]:
        issues: list[ListingIssue] = []

        def add(code: str, message: str, sku: str | None = None) -> None:
            issues.append(
                ListingIssue(
                    severity="ERROR",
                    code=code,
                    message=message,
                    design_id=product.design_id,
                    sku=sku,
                )
            )

        if not product.title_zh:
            add("MISSING_CHINESE_TITLE", "Temu SPU 商品名称尚未填写")
        if not product.asset_rights_confirmed:
            add("ASSET_RIGHTS_UNCONFIRMED", "图案及商品图片的商业使用权尚未确认")
        if not product.compliance_manifest or not product.compliance_manifest.is_file():
            add("MISSING_COMPLIANCE_REVIEW", "颜料 SDS 尚未关联到已审核的合规清单")
        else:
            try:
                raw_compliance = json.loads(product.compliance_manifest.read_text(encoding="utf-8"))
            except (json.JSONDecodeError, OSError) as exc:
                add("INVALID_COMPLIANCE_REVIEW", f"无法读取合规清单：{exc}")
            else:
                if not isinstance(raw_compliance, dict):
                    add("INVALID_COMPLIANCE_REVIEW", "合规清单必须是 JSON 对象")
                else:
                    missing_approvals: list[str] = []
                    if raw_compliance.get("jurisdiction") != "US":
                        missing_approvals.append("美国适用性")
                    if raw_compliance.get("product_linkage_confirmed") is not True:
                        missing_approvals.append("当前 SKU 与该颜料/SDS 的对应证明")
                    if raw_compliance.get("astm_d4236_toxicologist_review_confirmed") is not True:
                        missing_approvals.append("ASTM D-4236 毒理审核证明")
                    if raw_compliance.get("approved_for_publish") is not True:
                        missing_approvals.append("最终发布批准")
                    if missing_approvals:
                        add(
                            "COMPLIANCE_NOT_APPROVED",
                            "合规清单尚未批准发布，缺少：" + "、".join(missing_approvals),
                        )

        source_images = [product.main_image, *product.additional_images]
        for source_image in source_images:
            try:
                info = inspect_image(source_image)
            except ImageInspectionError as exc:
                add("INVALID_SOURCE_IMAGE", str(exc))
                continue
            if info.width_px <= min_dimension or info.height_px <= min_dimension:
                add(
                    "TEMU_IMAGE_DIMENSION_TOO_SMALL",
                    f"{info.path.name} 为 {info.width_px}x{info.height_px}px；模板要求宽高均大于 "
                    f"{min_dimension}px",
                )
            if info.width_px != info.height_px:
                add(
                    "TEMU_IMAGE_NOT_SQUARE",
                    f"{info.path.name} 不是 1:1 图片：{info.width_px}x{info.height_px}px",
                )
            if info.size_bytes >= max_bytes:
                add(
                    "TEMU_IMAGE_TOO_LARGE",
                    f"{info.path.name} 为 {info.size_bytes} bytes；模板要求小于 {max_bytes} bytes",
                )

        publish_refs = product.publish_image_refs
        if not min_carousel <= len(publish_refs) <= max_carousel:
            add(
                "INVALID_CAROUSEL_COUNT",
                f"轮播图素材 ID/URL 需要 {min_carousel}-{max_carousel} 个，当前为 "
                f"{len(publish_refs)} 个",
            )
        for ref in publish_refs:
            if not is_publishable_image_ref(ref):
                add(
                    "LOCAL_IMAGE_REF_NOT_UPLOADABLE",
                    f"模板只接受素材 ID 或可下载 URL，不能直接填写本地路径：{ref}",
                )

        colors = {sku.colors_count for sku in product.skus}
        frames = {sku.framed for sku in product.skus}
        if len(colors) > 1:
            add(
                "MIXED_SPU_PAINT_COUNTS",
                "颜料色数是 SPU 属性；不同色数必须拆成不同 SPU",
            )
        if len(frames) > 1:
            add(
                "MIXED_SPU_FRAME_TYPES",
                "框架类型是 SPU 属性；有框和无框必须拆成不同 SPU",
            )
        if "膏体" not in product.sensitive_attributes:
            add(
                "PASTE_SENSITIVE_ATTRIBUTE_UNCONFIRMED",
                "SDS 将颜料物态标为 Paste；正式发布前需确认敏感词属性是否填写“膏体”",
            )

        inventory_values: list[int] = []
        for sku in product.skus:
            if not sku.sku_image_ref or not is_publishable_image_ref(sku.sku_image_ref):
                add(
                    "MISSING_SKU_IMAGE_REF",
                    "SKU 预览图必须是店铺素材 ID 或有效 URL",
                    sku.factory_sku,
                )
            if sku.declared_price_cny is None:
                add("MISSING_DECLARED_PRICE", "缺少申报价格（CNY）", sku.factory_sku)
            if sku.warehouse_inventory is None:
                add("MISSING_WAREHOUSE_INVENTORY", "缺少领典仓库存", sku.factory_sku)
            else:
                inventory_values.append(sku.warehouse_inventory)
            if any(
                value is None
                for value in (
                    sku.package_length_cm,
                    sku.package_width_cm,
                    sku.package_height_cm,
                    sku.package_weight_g,
                )
            ):
                add(
                    "MISSING_PACKAGE_MEASUREMENT",
                    "缺少包装后三边尺寸或重量",
                    sku.factory_sku,
                )
        if inventory_values and max(inventory_values) <= 2:
            add("INSUFFICIENT_WAREHOUSE_INVENTORY", "至少一个 SKU 的领典仓库存必须大于 2")
        return issues

    def export(
        self,
        *,
        template_path: Path,
        mapping_path: Path,
        rows: list[CanonicalListingRow],
        output_path: Path,
        allow_test_fixture: bool = False,
    ) -> Path:
        mapping = self.validate_mapping(
            template_path, mapping_path, allow_test_fixture=allow_test_fixture
        )
        if not rows:
            raise TemuTemplateError("No valid listing rows are available for export")
        payload = {
            "mapping": mapping,
            "rows": [row.model_dump(mode="json") for row in rows],
        }
        output = output_path.expanduser().resolve()
        output.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(prefix="temu-template-") as temp_dir:
            payload_path = Path(temp_dir) / "payload.json"
            payload_path.write_text(
                json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
            )
            result = subprocess.run(
                [
                    self.node_executable,
                    str(self.writer_script),
                    str(template_path.expanduser().resolve()),
                    str(payload_path),
                    str(output),
                ],
                check=False,
                capture_output=True,
                text=True,
            )
        if result.returncode != 0 or not output.is_file():
            raise TemuTemplateError(
                "Temu template export failed: " + (result.stderr.strip() or result.stdout.strip())
            )
        self._restore_sheet_visibility(template_path.expanduser().resolve(), output)
        return output

    @staticmethod
    def _restore_sheet_visibility(template_path: Path, output_path: Path) -> None:
        """Restore metadata not currently round-tripped by Artifact Tool's XLSX importer."""
        source = load_workbook(template_path, read_only=True, data_only=False)
        target = load_workbook(output_path, read_only=False, data_only=False)
        try:
            states = {sheet.title: sheet.sheet_state for sheet in source.worksheets}
            for sheet in target.worksheets:
                if sheet.title in states:
                    sheet.sheet_state = states[sheet.title]
            target.save(output_path)
        finally:
            source.close()
            target.close()
